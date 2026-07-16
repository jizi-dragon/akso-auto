"""
v4 两步提取法：
第一遍：基于段落的章节提取（保留样式名称、隐式标题）。
第二遍：XML body 遍历 → 检测表格/图片 → 注入占位符到章节中。
"""
import argparse
import re
import os
from docx import Document

# ============================================================
# XML 命名空间 & 常量
# ============================================================
WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WPNS = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
ANS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

CHAPTER_STYLES = {'Title', 'Heading 1', 'Heading 2'}
ATTACHMENT_KEYWORDS = ('附件', '附表')
IMAGE_CAPTION_KEYWORDS = ('分布图', '取样点分布')
BODY_LEFT_INDENT = 360045
INDENT = '\u3000\u3000'

# ============================================================
# 辅助函数
# ============================================================

def get_section_num(title):
    """提取标题中的章节编号，如 '1.2.3'"""
    m = re.match(r'^(\d+(?:\.\d+)*)', title)
    return m.group(1) if m else None

def clean_title(title):
    """去除标题前的编号前缀"""
    return re.sub(r'^[\d.]+\s*', '', title)

def looks_like_implicit_heading(text, li):
    """判断一段文字是否像是隐式标题（无编号但语义上像标题）"""
    if len(text) > 25:
        return False
    if li == BODY_LEFT_INDENT:
        return False
    body_kw = ('备注', 'USL', 'LSL', '结论', '警戒限度', '行动限度',
               '小结', '从趋势图', '根据', '基于', '结合', '由于')
    for kw in body_kw:
        if text.startswith(kw):
            return False
    return True

def get_p_text(child):
    """从 XML 段落元素中提取文本"""
    text_nodes = child.findall('.//{' + WNS + '}t')
    return ''.join(n.text or '' for n in text_nodes).strip()

def has_image(child):
    """检查 XML 元素中是否包含图片"""
    drawings = child.findall('.//{' + WNS + '}drawing')
    inlines = child.findall('.//{' + WPNS + '}inline')
    blips = child.findall('.//{' + ANS + '}blip')
    return bool(drawings) or bool(inlines) or bool(blips)

def build_text_to_section_map(sections):
    """构建段落文本前60字符到章节索引的映射"""
    mapping = {}
    for sec_idx, sec in enumerate(sections):
        title = sec.get('title', '')
        clean = clean_title(title)
        key = title[:60]
        mapping[key] = sec_idx
        for _, ctext in sec.get('content', []):
            key = ctext[:60]
            mapping[key] = sec_idx
    return mapping

def find_section_for_media(prev_text, text_map):
    """通过前一段落文本查找表格/图片应归属的章节"""
    prev_key = prev_text[:60] if prev_text else ''
    if prev_key in text_map:
        return text_map[prev_key]
    for key, sec_idx in text_map.items():
        if prev_key and key and (prev_key.startswith(key[:40]) or key[:40].startswith(prev_key[:40])):
            return sec_idx
    for sec_idx, sec in enumerate(sections):
        title = clean_title(sec.get('title', ''))
        if title and prev_text and (title in prev_text or prev_text in title):
            return sec_idx
    return None

def infer_table_desc(section):
    """根据章节标题推断表格描述"""
    title = clean_title(section.get('title', ''))
    if '质量标准' in title: return '质量标准'
    if '监测频率' in title: return '监测频率表'
    if '取样点列表' in title: return '取样点列表'
    if '电导率' in title: return '电导率判定标准对照表'
    if '仪器仪表校验' in title: return '仪器仪表校验汇总表'
    if '维护保养' in title: return '维护保养汇总表'
    if '清洗' in title and '灭菌' in title: return '清洗灭菌钝化执行情况表'
    if '偏差' in title: return '偏差汇总表'
    if 'OOS' in title or 'OOT' in title: return 'OOS/OOT汇总表'
    if '变更' in title: return '变更汇总表'
    if '数据统计' in title or 'TOC' in title: return '数据统计表'
    return '数据表'

def infer_img_desc(section):
    """根据章节标题推断图片描述"""
    title = clean_title(section.get('title', ''))
    if '工艺流程' in title: return '工艺流程图'
    if '供应范围' in title or '分配系统' in title: return '分配系统流程图'
    if '趋势图' in title or '趋势分析' in title: return '趋势图'
    if '监测数据' in title:
        if '总有机碳' in title or 'TOC' in title: return '总有机碳趋势图'
        if '电导率' in title: return '电导率趋势图'
        if '微生物限度' in title: return '微生物限度趋势图'
        if '不挥发物' in title: return '不挥发物趋势图'
        if 'PH' in title: return 'PH趋势图'
        return '趋势图'
    if '运行' in title: return '运行数据图'
    return '图表'

# ============================================================
# 第一遍：基于段落的章节提取（源自 v3）
# ============================================================

def run_pass1(doc):
    """
    第一遍提取：逐段解析文档段落，识别章节标题、隐式标题和正文内容。
    返回 sections 列表。
    """
    raw_paras = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else ''
        li = para.paragraph_format.left_indent
        raw_paras.append({
            'text': text, 'style': style, 'left_indent': li,
            'has_num': bool(re.match(r'^\d+(?:\.\d+)*\s+', text))
        })

    sections = []
    current_section = None
    parent_chapter = None
    impl_counter = 0
    inside_attachment = False

    for pidx, p in enumerate(raw_paras):
        text = p['text']
        style = p['style']
        li = p['left_indent']

        if style.startswith('toc'):
            continue
        if '//' in text and len(text) < 20:
            continue
        if '有限公司' in text and len(sections) == 0:
            continue

        if any(kw in text for kw in IMAGE_CAPTION_KEYWORDS):
            if current_section:
                current_section.setdefault('media_placeholders', []).append(('image', text))
            continue

        if inside_attachment and not (style == 'Title' and get_section_num(text)):
            continue

        # Title / Heading 样式
        if style in CHAPTER_STYLES or style == 'Title':
            sec_num = get_section_num(text)
            if not sec_num:
                continue
            is_att = any(kw in text for kw in ATTACHMENT_KEYWORDS)
            if '.' in sec_num:
                sections.append({
                    'type': 'sub', 'title': text, 'sec_num': sec_num,
                    'content': [], 'media_placeholders': [],
                    'is_attachment': is_att
                })
                current_section = sections[-1]
                if is_att:
                    inside_attachment = True
                continue
            sections.append({
                'type': 'chapter', 'title': text, 'sec_num': sec_num,
                'content': [], 'media_placeholders': [],
                'is_attachment': is_att
            })
            current_section = sections[-1]
            parent_chapter = sections[-1]
            impl_counter = 0
            inside_attachment = is_att
            continue

        # 带显示编号的 Normal 段落
        if p['has_num']:
            sec_num = get_section_num(text)
            if sec_num and '.' in sec_num:
                sections.append({
                    'type': 'sub', 'title': text, 'sec_num': sec_num,
                    'content': [], 'media_placeholders': []
                })
                current_section = sections[-1]
            elif sec_num and parent_chapter and parent_chapter.get('is_attachment'):
                continue
            continue

        # 隐式标题？
        if looks_like_implicit_heading(text, li):
            is_followed_by_body = False
            for j in range(pidx + 1, min(pidx + 5, len(raw_paras))):
                nxt = raw_paras[j]
                if nxt.get('has_num') or nxt.get('style') in CHAPTER_STYLES or nxt.get('style') == 'Title':
                    break
                if nxt.get('left_indent') == BODY_LEFT_INDENT:
                    is_followed_by_body = True
                    break
            if is_followed_by_body and parent_chapter:
                impl_counter += 1
                base_num = parent_chapter['sec_num']
                impl_num = f"{base_num}.{impl_counter}"
                sections.append({
                    'type': 'sub', 'title': f"{impl_num} {text}", 'sec_num': impl_num,
                    'content': [], 'media_placeholders': [], 'is_implicit': True
                })
                current_section = sections[-1]
                continue

        # 正文
        if current_section and not current_section.get('is_attachment'):
            current_section.setdefault('content', []).append(('text', text))

    return sections

# ============================================================
# 第二遍：XML body 遍历 → 检测表格/图片 → 注入到章节
# ============================================================

def run_pass2(doc, sections):
    """
    第二遍提取：遍历 XML body，检测表格和图片元素，
    根据上下文将其作为占位符注入到对应章节中。
    直接修改传入的 sections 列表。
    """
    text_map = build_text_to_section_map(sections)
    media_with_context = []

    last_para_text = ''
    for child in doc.element.body:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag == 'p':
            text = get_p_text(child)
            is_img = has_image(child)

            if is_img:
                media_with_context.append({'type': 'image', 'prev_text': last_para_text})
                if text:
                    last_para_text = text
            elif text:
                last_para_text = text
            continue

        if tag == 'tbl':
            media_with_context.append({'type': 'table', 'prev_text': last_para_text})
            continue

        if tag == 'sdt':
            tbls = child.findall('.//{' + WNS + '}tbl')
            if tbls:
                media_with_context.append({'type': 'table', 'prev_text': last_para_text})
            continue

    # 将媒体元素注入到对应章节
    for media in media_with_context:
        sec_idx = find_section_for_media(media['prev_text'], text_map)
        if sec_idx is not None and sec_idx < len(sections):
            sec = sections[sec_idx]
            if media['type'] == 'table':
                desc = infer_table_desc(sec)
                sec.setdefault('media_placeholders', []).append(('table', desc))
            else:
                desc = infer_img_desc(sec)
                sec.setdefault('media_placeholders', []).append(('image', desc))

# ============================================================
# 格式化输出
# ============================================================

def format_output(doc_path, sections):
    """
    将提取的章节结构格式化为三部分输出：
    1. [TOC 目录]
    2. [完整子标题结构]
    3. [大纲 + 文字段落]
    返回字符串。
    """
    doc_title = os.path.splitext(os.path.basename(doc_path))[0]
    lines = []
    lines.append('=' * 60)
    lines.append(doc_title)
    lines.append('=' * 60)
    lines.append('')

    # 第一部分：TOC 目录
    lines.append('[TOC 目录]')
    lines.append('-' * 40)
    for sec in sections:
        if sec['type'] == 'chapter':
            lines.append(f"{sec['sec_num']} {clean_title(sec['title'])}")
    lines.append('')

    # 第二部分：完整子标题结构
    lines.append('[完整子标题结构]')
    lines.append('-' * 40)
    for sec in sections:
        sc = sec['sec_num']
        clean = clean_title(sec['title'])
        if sec['type'] == 'chapter':
            lines.append(f'{sc} {clean}')
        else:
            depth = sc.count('.')
            lines.append(f"{'  ' * depth}{sc} {clean}")
    lines.append('')

    # 第三部分：大纲 + 文字段落
    lines.append('[大纲 + 文字段落]')
    lines.append('-' * 40)
    for sec in sections:
        sc = sec['sec_num']
        clean = clean_title(sec['title'])

        if sec['type'] == 'chapter':
            lines.append(f'--- Ch{sc} ---')

        if sec.get('is_attachment'):
            lines.append(f'{sc} {clean}')
            for mtype, mdesc in sec.get('media_placeholders', []):
                tag = '[表格：' if mtype == 'table' else '[图片：'
                lines.append(f'{tag}{mdesc}]')
            if not sec.get('media_placeholders'):
                lines.append(f'{INDENT}[以下为附件数据表格，不含文字段落]')
            lines.append('')
            continue

        lines.append(f'{sc} {clean}')

        for mtype, mdesc in sec.get('media_placeholders', []):
            tag = '[表格：' if mtype == 'table' else '[图片：'
            lines.append(f'{tag}{mdesc}]')

        for _, ctext in sec.get('content', []):
            lines.append(f'{INDENT}{ctext}')

        lines.append('')

    return '\n'.join(lines)

# ============================================================
# 打印摘要统计
# ============================================================

def print_summary(sections):
    """打印提取摘要统计信息"""
    print(f'章节总数: {len(sections)}')
    for sec in sections:
        tc = len(sec.get('content', []))
        mp = len(sec.get('media_placeholders', []))
        imp = '[隐式]' if sec.get('is_implicit') else ''
        att = '[附件]' if sec.get('is_attachment') else ''
        mi = ','.join([f'{m[0][:1]}:{m[1][:15]}' for m in sec.get('media_placeholders', [])])
        print(f"  {sec['type']:7} {imp:6} {att:6} {sec['sec_num']:8} | {clean_title(sec['title'])} (正文={tc} 媒体=[{mi}])")

# ============================================================
# 主入口
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description='从 .docx 文件中提取文档大纲结构，包含目录、子标题层级和正文段落'
    )
    parser.add_argument(
        '--input', required=True,
        help='输入的 .docx 文件路径'
    )
    parser.add_argument(
        '--output', required=True,
        help='输出的 .txt 文件路径'
    )
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f'错误：输入文件不存在 —— {args.input}')
        return

    print(f'正在读取文档: {args.input}')
    doc = Document(args.input)

    print('第一遍提取：基于段落解析章节结构...')
    sections = run_pass1(doc)

    print('第二遍提取：XML body 遍历，检测表格和图片...')
    run_pass2(doc, sections)

    print('格式化输出...')
    output = format_output(args.input, sections)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    with open(args.output, 'w', encoding='utf-8') as f:
        f.write(output)

    print(f'输出文件已保存: {args.output}')
    print()
    print_summary(sections)

if __name__ == '__main__':
    main()
